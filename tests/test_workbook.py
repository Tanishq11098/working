import os
import pytest
import openpyxl
from bot.storage import (
    ExcelStorageManager,
    SHEET_ALL_LISTINGS,
    SHEET_SOURCE_STATS,
    SHEET_SEEN_LISTINGS,
    SHEET_TRACKER
)
from bot.models import JobListing, SourceExecutionStats

def test_workbook_initialization(tmp_path):
    wb_path = os.path.join(tmp_path, "internships.xlsx")
    manager = ExcelStorageManager(filepath=wb_path)

    # File must exist
    assert os.path.exists(wb_path)
    
    # Must contain exactly the 4 required sheets
    wb = openpyxl.load_workbook(wb_path)
    assert set(wb.sheetnames) == {
        SHEET_ALL_LISTINGS,
        SHEET_SOURCE_STATS,
        SHEET_SEEN_LISTINGS,
        SHEET_TRACKER
    }
    wb.close()

def test_workbook_data_persistence(tmp_path):
    wb_path = os.path.join(tmp_path, "internships.xlsx")
    manager = ExcelStorageManager(filepath=wb_path)

    # Initial insertion of a listing
    listing = JobListing(
        job_id="TEST-101",
        title="Python Developer Intern",
        company="TechCorp",
        location="Noida",
        match_score=85,
        match_reasons=["+20 pts: NCR location", "+30 pts: Python keyword"],
        original_url="https://example.com/job/101?ref=test",
        application_url="https://example.com/job/101"
    )
    stat = SourceExecutionStats(
        run_id="RUN-1",
        run_datetime_ist="2026-10-18 10:00:00",
        source_name="TestSource",
        records_retrieved=1,
        new_listings=1,
        listings_passing_filters=1
    )

    success = manager.update_workbook([listing], [stat], tracker_threshold=60)
    assert success is True

    # Re-open manager and verify deduplication
    manager2 = ExcelStorageManager(filepath=wb_path)
    assert manager2.is_duplicate(listing) is True

    # Check Tracker sheet contains listing
    wb = openpyxl.load_workbook(wb_path)
    ws_tracker = wb[SHEET_TRACKER]
    rows = list(ws_tracker.iter_rows(values_only=True))
    assert len(rows) == 2  # Header + 1 record
    assert rows[1][0] == "TEST-101"
    assert rows[1][6] == "Not Applied"

    # Verify seen_listings
    ws_seen = wb[SHEET_SEEN_LISTINGS]
    seen_rows = list(ws_seen.iter_rows(values_only=True))
    assert len(seen_rows) == 2
    assert seen_rows[1][0] == "TEST-101"
    wb.close()

def test_workbook_validation_detects_missing_sheet(tmp_path):
    wb_path = os.path.join(tmp_path, "internships.xlsx")
    manager = ExcelStorageManager(filepath=wb_path)

    # Corrupt workbook by deleting a sheet
    wb = openpyxl.load_workbook(wb_path)
    del wb[SHEET_TRACKER]
    wb.save(wb_path)
    wb.close()

    valid, reason = manager.validate_workbook()
    assert valid is False
    assert "Tracker" in reason
